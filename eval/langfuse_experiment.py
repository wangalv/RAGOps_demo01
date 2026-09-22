"""
M5.3 — Langfuse Experiments (v4 API)

Uses langfuse.run_experiment() to create proper Experiments visible in the
Langfuse UI under Datasets → NSW-Health-Services-Act-Golden-Set → Experiments.

Each retriever technique becomes one named experiment run, with R@3/R@5/P@3/P@5
scores automatically computed and stored per dataset item.

Run from LangGraphLearning/:
    .venv/bin/python eval/langfuse_experiment.py          # all retrievers
    .venv/bin/python eval/langfuse_experiment.py --fast   # skip LLM-based
"""

import argparse
import os
import re
import sys
sys.path.insert(0, ".")

import openpyxl
from dotenv import load_dotenv
load_dotenv()

from openai import OpenAI as _OpenAI
from langfuse import Langfuse

from src.rag.chunker import chunk_by_section, chunk_parent_child, chunk_by_subsection, chunk_subsection_parent_child
from src.rag.vector_store import VectorStore
from src.rag.reranker import Reranker
from src.rag.hybrid_retriever import HybridRetriever
from src.rag.hyde_retriever import HyDERetriever, LegalHyDERetriever, generate_hypothetical_doc
from src.rag.legal_query_expansion_retriever import LegalQueryExpansionRetriever, expand_query
from src.rag.routed_retriever import RoutedRetriever
from src.rag.multi_query_retriever import MultiQueryRetriever, generate_query_variants
from src.rag.stepback_retriever import StepBackRetriever, generate_stepback_query
from src.rag.ensemble_retriever import EnsembleRetriever
from src.rag.query_router import classify_query
from src.rag.llm_reranker import llm_rerank_legal

ACT_TEXT_FILE     = "tests/rag/health_services_act.txt"
GOLDEN_SET_FILE   = "GoldenSet_v2.xlsx"
SIM_LABELED_FILE      = "GoldenSet_from_100_Simulated_Queries.xlsx"
PARENT_COLLECTION     = "health-services-act-baseline"
CHILD_COLLECTION      = "health-services-act-children"
SUBSEC_COLLECTION     = "health-services-act-subsections"
SUBSEC_CHILD_COLLECTION = "health-services-act-subsec-children"
DATASET_NAME          = "NSW-Health-Services-Act-Golden-Set"
SIM_DATASET_NAME      = "NSW-HealthAct-SimulatedQueries-Labeled-100"
SIM_GROUNDED_NAME     = "NSW-HealthAct-SimulatedQueries-Grounded-77"
CANDIDATE_N       = 20
EVAL_K            = [3, 5]

langfuse = Langfuse(
    public_key=os.environ["LANGFUSE_PUBLIC_KEY"],
    secret_key=os.environ["LANGFUSE_SECRET_KEY"],
    host=os.environ.get("LANGFUSE_HOST", "https://cloud.langfuse.com"),
)

_llm = _OpenAI(
    base_url=os.environ["Z_AI_BASE_URL"],
    api_key=os.environ["Z_AI_API_KEY"],
)


def llm_rerank(query: str, chunks: list[dict], top_k: int) -> list[dict]:
    """Ask LLM to rerank chunks by relevance, return top_k."""
    if not chunks:
        return []
    chunk_list = "\n\n".join(
        f"[{i+1}] {c['text'][:500]}" for i, c in enumerate(chunks)
    )
    prompt = (
        f"You are evaluating passages from the NSW Health Services Act 1997.\n\n"
        f"Query: {query}\n\n"
        f"Passages:\n{chunk_list}\n\n"
        f"Rank these {len(chunks)} passages from MOST to LEAST relevant to the query.\n"
        f"Reply with ONLY a comma-separated list of numbers, e.g.: 3,1,5,2,4\n"
        f"No explanation, no other text."
    )
    try:
        resp = _llm.chat.completions.create(
            model=os.environ["Z_AI_CHAT_MODEL"],
            messages=[{"role": "user", "content": prompt}],
            temperature=0,
            max_tokens=50,
        )
        raw = (resp.choices[0].message.content or "").strip()
        seen: set[int] = set()
        ranking = []
        for x in raw.split(","):
            try:
                idx = int(x.strip()) - 1
                if 0 <= idx < len(chunks) and idx not in seen:
                    seen.add(idx)
                    ranking.append(idx)
            except ValueError:
                pass
        # append any missed indices (fallback)
        for i in range(len(chunks)):
            if i not in seen:
                ranking.append(i)
        return [chunks[i] for i in ranking[:top_k]]
    except Exception:
        return chunks[:top_k]


# ── Golden Set ────────────────────────────────────────────────────────────────

def parse_expected_sections(raw) -> list[str]:
    """Parse 'Expected Section(s)' field → list of section ids like 's.17'."""
    if not raw:
        return []
    raw_str = str(raw).lower()
    if any(p in raw_str for p in ["no direct", "no general", "dictionary"]):
        return []
    sections = []
    for part in str(raw).replace("–", "-").replace("—", "-").split(";"):
        part = part.strip()
        m = re.match(r"s\s+(\d+[A-Za-z]*)", part, re.IGNORECASE)
        if m:
            sections.append(f"s.{m.group(1)}")
    return sections


def load_simulated_labeled_set(path: str) -> list[dict]:
    """Load 100 simulated queries with human-labeled expected sections."""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Golden Review"]
    rows = list(ws.iter_rows(values_only=True))
    header = rows[0]
    col = {name: i for i, name in enumerate(header)}
    cases = []
    for row in rows[1:]:
        if not row[col["ID"]]:
            continue
        cases.append({
            "id":         row[col["ID"]],
            "query":      row[col["Query"]],
            "relevant":   parse_expected_sections(row[col["Expected Section(s)"]]),
            "query_type": row[col["Original Query Type"]],
            "difficulty": row[col["Grounded"]],
            "candidate":  row[col["Golden Set Candidate"]] or "",
        })
    return cases


def setup_sim_dataset(cases: list[dict], dataset_name: str = SIM_DATASET_NAME, id_prefix: str = "simlabeled"):
    try:
        langfuse.create_dataset(
            name=dataset_name,
            description=f"{len(cases)} simulated queries with human-labeled sections",
        )
        print(f"  Created dataset '{dataset_name}'")
    except Exception:
        print(f"  Dataset '{dataset_name}' already exists, reusing")

    for case in cases:
        try:
            langfuse.create_dataset_item(
                dataset_name=dataset_name,
                input={"query": case["query"]},
                expected_output={"relevant_sections": case["relevant"]},
                metadata={
                    "case_id":    case["id"],
                    "query_type": case["query_type"],
                    "grounded":   case["difficulty"],
                    "candidate":  case["candidate"],
                },
                id=f"{id_prefix}-{case['id']}",
            )
        except Exception:
            pass  # item already exists, skip

    dataset = langfuse.get_dataset(dataset_name)
    print(f"  {len(dataset.items)} items ready\n")
    return dataset


def load_golden_set(path: str) -> list[dict]:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Golden Set v2"]
    rows = list(ws.iter_rows(values_only=True))
    header = rows[0]
    col = {name: i for i, name in enumerate(header)}
    cases = []
    for row in rows[1:]:
        if not row[col["ID"]]:
            continue
        primary   = row[col["Primary Section"]] or ""
        secondary = row[col["Secondary Sections"]] or ""
        relevant  = []
        for ref in (primary + "," + secondary).split(","):
            ref = ref.strip()
            if ref and ref.lower() != "none":
                relevant.append(ref)
        cases.append({
            "id":         row[col["ID"]],
            "query":      row[col["Question"]],
            "relevant":   relevant,
            "query_type": row[col["Query Type"]],
            "difficulty": row[col["Difficulty"]],
        })
    return cases


# ── Metrics ───────────────────────────────────────────────────────────────────

def extract_section_prefix(chunk_id: str) -> str:
    m = re.match(r"^(s\.[^_]+)", chunk_id)
    return m.group(1) if m else chunk_id


def recall_at_k(retrieved_sections: list[str], relevant: list[str], k: int) -> float:
    if not relevant:
        return 1.0
    found = len({s for s in relevant if s in set(retrieved_sections[:k])})
    return found / len(relevant)


def precision_at_k(retrieved_sections: list[str], relevant: list[str], k: int) -> float:
    if k == 0:
        return 0.0
    found = len({s for s in relevant if s in set(retrieved_sections[:k])})
    return found / k


# ── Dataset setup ─────────────────────────────────────────────────────────────

def setup_dataset(cases: list[dict]):
    try:
        langfuse.create_dataset(
            name=DATASET_NAME,
            description=f"NSW Health Services Act — {len(cases)} evaluation queries",
        )
        print(f"  Created dataset '{DATASET_NAME}'")
    except Exception:
        print(f"  Dataset '{DATASET_NAME}' already exists, reusing")

    for case in cases:
        langfuse.create_dataset_item(
            dataset_name=DATASET_NAME,
            input={"query": case["query"]},
            expected_output={"relevant_sections": case["relevant"]},
            metadata={
                "case_id":    case["id"],
                "query_type": case["query_type"],
                "difficulty": case["difficulty"],
            },
            id=f"item-{case['id']}",
        )

    dataset = langfuse.get_dataset(DATASET_NAME)
    print(f"  {len(dataset.items)} items ready\n")
    return dataset


# ── Experiment runner ─────────────────────────────────────────────────────────

def run_one_experiment(retriever_name: str, retriever_fn, dataset, ks: list[int], dataset_name: str = DATASET_NAME):
    """Run one retriever as a Langfuse Experiment (v4)."""
    print(f"  Running: {retriever_name}")

    def task(item):
        query = item.input["query"]
        retrieved = retriever_fn(query)
        sections = [extract_section_prefix(r["chunk_id"]) for r in retrieved]
        return {"retrieved_sections": sections}

    def make_evaluator(k, metric):
        def evaluator(*, input, output, expected_output, metadata=None, **kwargs):
            retrieved = (output or {}).get("retrieved_sections", [])
            relevant  = (expected_output or {}).get("relevant_sections", [])
            if metric == "recall":
                value = recall_at_k(retrieved, relevant, k)
            else:
                value = precision_at_k(retrieved, relevant, k)
            return {"name": f"{metric.upper()}@{k}", "value": value}
        return evaluator

    evaluators = []
    for k in ks:
        evaluators.append(make_evaluator(k, "recall"))
        evaluators.append(make_evaluator(k, "precision"))

    result = langfuse.run_experiment(
        name=dataset_name,
        run_name=retriever_name,
        data=dataset.items,
        task=task,
        evaluators=evaluators,
        max_concurrency=1,
    )

    # Collect all metric values from item_results
    metric_vals: dict[str, list[float]] = {}
    for ir in result.item_results:
        for ev in ir.evaluations:
            metric_vals.setdefault(ev.name, []).append(ev.value)
    n = len(result.item_results)
    expected_metrics = [f"{m}@{k}" for k in ks for m in ["RECALL", "PRECISION"]]
    for metric in expected_metrics:
        vals = metric_vals.get(metric, [])
        avg = sum(vals) / len(vals) if vals else 0
        print(f"    {metric:<14} = {avg:.3f}  (n={len(vals)})")
    if result.dataset_run_url:
        print(f"    → {result.dataset_run_url}")


# ── Main ──────────────────────────────────────────────────────────────────────

def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--fast", action="store_true",
                        help="Skip LLM-based techniques")
    parser.add_argument("--only-dense", action="store_true",
                        help="Only run Dense@3 and Dense@5")
    parser.add_argument("--run", type=str, default=None,
                        help="Comma-separated list of retriever names to run, e.g. 'Dense@3,Dense@5'")
    parser.add_argument("--sim-labeled", action="store_true",
                        help="Use 100 simulated+human-labeled queries instead of original golden set")
    parser.add_argument("--grounded-only", action="store_true",
                        help="With --sim-labeled: exclude unanswerable queries (relevant=[]), keeps ~77 grounded")
    parser.add_argument("--limit", type=int, default=None,
                        help="Only run the first N queries (for quick smoke tests)")
    parser.add_argument("--eval-k", type=str, default=None,
                        help="Comma-separated k values for recall/precision, e.g. '3,5,20' (default: 3,5)")
    args = parser.parse_args()

    if args.sim_labeled:
        print("Loading Simulated Labeled Set (100 queries)...")
        cases = load_simulated_labeled_set(SIM_LABELED_FILE)
        print(f"  {len(cases)} queries  (Yes={sum(1 for c in cases if c['candidate']=='Yes')}  "
              f"Rewrite={sum(1 for c in cases if 'Rewrite' in c['candidate'])}  "
              f"No={sum(1 for c in cases if c['candidate']=='No')})")
        if args.grounded_only:
            before = len(cases)
            cases = [c for c in cases if c["relevant"]]
            print(f"  [--grounded-only] Filtered {before - len(cases)} unanswerable → {len(cases)} grounded queries")
        print()
    else:
        print("Loading Golden Set...")
        cases = load_golden_set(GOLDEN_SET_FILE)
        print(f"  {len(cases)} queries\n")

    if args.limit:
        cases = cases[: args.limit]
        print(f"  [--limit] Using first {len(cases)} queries only\n")

    eval_ks = [int(k) for k in args.eval_k.split(",")] if args.eval_k else EVAL_K

    print("Loading Act and building indexes...")
    with open(ACT_TEXT_FILE, encoding="utf-8") as f:
        act_text = f.read()

    parents, children = chunk_parent_child(act_text)
    chunks = chunk_by_section(act_text)
    subsec_chunks, subsec_children = chunk_subsection_parent_child(act_text)

    parent_store = VectorStore(collection_name=PARENT_COLLECTION, persist_dir="./chroma_db")
    if parent_store.count() != len(chunks):
        parent_store.reset()
        parent_store.add_chunks(chunks)
    else:
        print(f"  Parent store: {parent_store.count()} chunks (reused)")

    child_store = VectorStore(collection_name=CHILD_COLLECTION, persist_dir="./chroma_db")
    if child_store.count() != len(children):
        child_store.reset()
        child_store.add_children(children)
    else:
        print(f"  Child store:  {child_store.count()} chunks (reused)")

    subsec_store = VectorStore(collection_name=SUBSEC_COLLECTION, persist_dir="./chroma_db")
    if subsec_store.count() != len(subsec_chunks):
        subsec_store.reset()
        subsec_store.add_chunks(subsec_chunks)
    else:
        print(f"  Subsec store: {subsec_store.count()} chunks (reused)")

    subsec_child_store = VectorStore(collection_name=SUBSEC_CHILD_COLLECTION, persist_dir="./chroma_db")
    if subsec_child_store.count() != len(subsec_children):
        subsec_child_store.reset()
        subsec_child_store.add_children(subsec_children)
    else:
        print(f"  SubsecChild:  {subsec_child_store.count()} chunks (reused)")

    reranker = Reranker()
    hybrid   = HybridRetriever(vector_store=parent_store)
    hybrid.index(chunks)
    hyde       = HyDERetriever(vector_store=parent_store)
    legal_hyde = LegalHyDERetriever(vector_store=parent_store)
    subsec_hyde = HyDERetriever(vector_store=subsec_store)
    lqe       = LegalQueryExpansionRetriever(hybrid_retriever=hybrid)
    mq        = MultiQueryRetriever(vector_store=parent_store, reranker=reranker)
    sb        = StepBackRetriever(vector_store=parent_store, reranker=reranker)
    routed    = RoutedRetriever(lqe=lqe, stepback=sb, multiquery=mq, hyde=hyde,
                                reranker=reranker, candidate_n=CANDIDATE_N)

    subsec_hybrid = HybridRetriever(vector_store=subsec_store)
    subsec_hybrid.index(subsec_chunks)
    subsec_lqe = LegalQueryExpansionRetriever(hybrid_retriever=subsec_hybrid)
    subsec_mq  = MultiQueryRetriever(vector_store=subsec_store, reranker=reranker)
    subsec_sb  = StepBackRetriever(vector_store=subsec_store, reranker=reranker)
    subsec_routed = RoutedRetriever(lqe=subsec_lqe, stepback=subsec_sb, multiquery=subsec_mq,
                                    hyde=subsec_hyde, reranker=reranker, candidate_n=CANDIDATE_N)

    retrievers = {
        "Dense@3":           lambda q: parent_store.search(q, top_k=3),
        "Dense@5":           lambda q: parent_store.search(q, top_k=5),
    }

    if not args.only_dense:
        retrievers.update({
            "Dense→Reranked@3":  lambda q: reranker.rerank(q, parent_store.search(q, top_k=CANDIDATE_N), top_k=3),
            "Dense→Reranked@5":  lambda q: reranker.rerank(q, parent_store.search(q, top_k=CANDIDATE_N), top_k=5),
            "Hybrid→Reranked@3": lambda q: reranker.rerank(q, hybrid.search(q, top_k=CANDIDATE_N), top_k=3),
            "Hybrid→Reranked@5": lambda q: reranker.rerank(q, hybrid.search(q, top_k=CANDIDATE_N), top_k=5),
            "ParentChild→Rnk@3": lambda q: reranker.rerank(q, child_store.get_parents_by_ids(parent_store, [c["parent_id"] for c in child_store.search_children(q, top_k=CANDIDATE_N)]), top_k=3),
            "ParentChild→Rnk@5": lambda q: reranker.rerank(q, child_store.get_parents_by_ids(parent_store, [c["parent_id"] for c in child_store.search_children(q, top_k=CANDIDATE_N)]), top_k=5),
        })

    if not args.fast and not args.only_dense:
        retrievers.update({
            "HyDE→Reranked@3":       lambda q: reranker.rerank(q, hyde.search(q, top_k=CANDIDATE_N)[0], top_k=3),
            "HyDE→Reranked@5":       lambda q: reranker.rerank(q, hyde.search(q, top_k=CANDIDATE_N)[0], top_k=5),
            "LegalHyDE→Reranked@3":  lambda q: reranker.rerank(q, legal_hyde.search(q, top_k=CANDIDATE_N)[0], top_k=3),
            "LegalHyDE→Reranked@5":  lambda q: reranker.rerank(q, legal_hyde.search(q, top_k=CANDIDATE_N)[0], top_k=5),
            "LQE→Hybrid→Reranked@3": lambda q: reranker.rerank(q, lqe.search(q, top_k=CANDIDATE_N)[0], top_k=3),
            "LQE→Hybrid→Reranked@5": lambda q: reranker.rerank(q, lqe.search(q, top_k=CANDIDATE_N)[0], top_k=5),
            "Routed→Reranked@3":     lambda q: routed.search(q, top_k=3)[0],
            "Routed→Reranked@5":     lambda q: routed.search(q, top_k=5)[0],
            "HyDE→LLMRerank@3": lambda q: llm_rerank(q, reranker.rerank(q, hyde.search(q, top_k=CANDIDATE_N)[0], top_k=5), top_k=3),
            "HyDE→LLMRerank@5": lambda q: llm_rerank(q, reranker.rerank(q, hyde.search(q, top_k=CANDIDATE_N)[0], top_k=5), top_k=5),
            "MultiQuery→Rnk@3": lambda q: mq.search(q, n_variants=3, per_query_k=CANDIDATE_N, final_k=3)[0],
            "MultiQuery→Rnk@5": lambda q: mq.search(q, n_variants=3, per_query_k=CANDIDATE_N, final_k=5)[0],
            "StepBack→Rnk@3":   lambda q: sb.search(q, per_query_k=CANDIDATE_N, final_k=3)[0],
            "StepBack→Rnk@5":   lambda q: sb.search(q, per_query_k=CANDIDATE_N, final_k=5)[0],
            "SubsecHyDE→Rnk@3":    lambda q: reranker.rerank(q, subsec_hyde.search(q, top_k=CANDIDATE_N)[0], top_k=3),
            "SubsecHyDE→Rnk@5":    lambda q: reranker.rerank(q, subsec_hyde.search(q, top_k=CANDIDATE_N)[0], top_k=5),
            "SubsecPC→Rnk@3":      lambda q: reranker.rerank(q, subsec_child_store.get_parents_by_ids(subsec_store, [c["parent_id"] for c in subsec_child_store.search_children(q, top_k=CANDIDATE_N)]), top_k=3),
            "SubsecPC→Rnk@5":      lambda q: reranker.rerank(q, subsec_child_store.get_parents_by_ids(subsec_store, [c["parent_id"] for c in subsec_child_store.search_children(q, top_k=CANDIDATE_N)]), top_k=5),
            "SubsecRouted→Rnk@5":  lambda q: subsec_routed.search(q, top_k=5)[0],
        })

        def _subsec_routed_pc(q: str, top_k: int = 5) -> list[dict]:
            """Routing + parent-child at subsection level + rerank."""
            qtype = classify_query(q)
            seen: set[str] = set()
            children: list[dict] = []

            def _collect(search_q: str) -> None:
                for c in subsec_child_store.search_children(search_q, top_k=CANDIDATE_N):
                    if c["chunk_id"] not in seen:
                        seen.add(c["chunk_id"])
                        children.append(c)

            if qtype == "Penalty":
                _collect(expand_query(q))
            elif qtype == "Prohibition":
                _collect(q)
                _collect(generate_stepback_query(q))
            elif qtype == "CrossSection":
                for sq in [q] + generate_query_variants(q, n=3):
                    _collect(sq)
            else:
                _collect(generate_hypothetical_doc(q))

            parents = subsec_child_store.get_parents_by_ids(
                subsec_store, [c["parent_id"] for c in children]
            )
            return reranker.rerank(q, parents, top_k=top_k)

        def _subsec_routed_pc_cascade(q: str) -> list[dict]:
            """Same as SubsecRoutedPC but uses cascade reranking (MiniLM→top10→BGE→top5)."""
            qtype = classify_query(q)
            seen: set[str] = set()
            children: list[dict] = []

            def _collect(search_q: str) -> None:
                for c in subsec_child_store.search_children(search_q, top_k=CANDIDATE_N):
                    if c["chunk_id"] not in seen:
                        seen.add(c["chunk_id"])
                        children.append(c)

            if qtype == "Penalty":
                _collect(expand_query(q))
            elif qtype == "Prohibition":
                _collect(q)
                _collect(generate_stepback_query(q))
            elif qtype == "CrossSection":
                for sq in [q] + generate_query_variants(q, n=3):
                    _collect(sq)
            else:
                _collect(generate_hypothetical_doc(q))

            parents = subsec_child_store.get_parents_by_ids(
                subsec_store, [c["parent_id"] for c in children]
            )
            return reranker.cascade_rerank(q, parents, mid_k=10, top_k=5)

        def _subsec_routed_pc_llm(q: str) -> list[dict]:
            """SubsecRoutedPC top-10 via MiniLM, then LLM legal reranker → top-5."""
            top10 = _subsec_routed_pc(q, top_k=10)
            return llm_rerank_legal(q, top10, top_k=5)

        retrievers.update({
            "SubsecRoutedPC→Rnk@5":      _subsec_routed_pc,
            "SubsecRoutedPC→Rnk@10":     lambda q: _subsec_routed_pc(q, top_k=10),
            "SubsecRoutedPC→Rnk@20":     lambda q: _subsec_routed_pc(q, top_k=20),
            "SubsecRoutedPC→Cascade@5":  _subsec_routed_pc_cascade,
            "SubsecRoutedPC→LLMRnk@5":   _subsec_routed_pc_llm,
        })

        # Ensemble retrievers — merge candidate pools then rerank
        ensemble_dh = EnsembleRetriever(
            retrievers=[
                {"name": "Dense",  "fn": lambda q, n: parent_store.search(q, top_k=n)},
                {"name": "HyDE",   "fn": lambda q, n: hyde.search(q, top_k=n)[0]},
            ],
            reranker=reranker,
            candidate_n=CANDIDATE_N,
        )
        ensemble_all = EnsembleRetriever(
            retrievers=[
                {"name": "Dense",       "fn": lambda q, n: parent_store.search(q, top_k=n)},
                {"name": "HyDE",        "fn": lambda q, n: hyde.search(q, top_k=n)[0]},
                {"name": "MultiQuery",  "fn": lambda q, n: mq.search(q, n_variants=3, per_query_k=n, final_k=n)[0]},
                {"name": "StepBack",    "fn": lambda q, n: sb.search(q, per_query_k=n, final_k=n)[0]},
            ],
            reranker=reranker,
            candidate_n=CANDIDATE_N,
        )
        retrievers.update({
            "Ensemble-DH@3":  lambda q: ensemble_dh.search(q, final_k=3),
            "Ensemble-DH@5":  lambda q: ensemble_dh.search(q, final_k=5),
            "Ensemble-All@3": lambda q: ensemble_all.search(q, final_k=3),
            "Ensemble-All@5": lambda q: ensemble_all.search(q, final_k=5),
        })
    else:
        print("  [--fast] Skipping HyDE, MultiQuery, Step-back\n")

    print("Setting up Langfuse dataset...")
    if args.sim_labeled:
        if args.grounded_only:
            ds_name, id_prefix = SIM_GROUNDED_NAME, "grounded"
        else:
            ds_name, id_prefix = SIM_DATASET_NAME, "simlabeled"
        dataset = setup_sim_dataset(cases, ds_name, id_prefix)
        target_dataset_name = ds_name
    else:
        dataset = setup_dataset(cases)
        target_dataset_name = DATASET_NAME

    if args.run:
        only = {n.strip() for n in args.run.split(",")}
        retrievers = {k: v for k, v in retrievers.items() if k in only}
        print(f"  [--run] Only: {', '.join(retrievers)}\n")

    print("Running experiments...\n")
    for name, fn in retrievers.items():
        run_one_experiment(name, fn, dataset, eval_ks, dataset_name=target_dataset_name)

    langfuse.flush()
    print(f"\n✓ Done — view at: https://cloud.langfuse.com")
    print(f"  Datasets → {target_dataset_name} → Experiments")


if __name__ == "__main__":
    main()
